from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT.parent / 'APEXFORMZ V2 Professional PRD.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(.65)
sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.75)
sec.right_margin = Inches(.75)
sec.header_distance = Inches(.25)
sec.footer_distance = Inches(.25)
sec.different_first_page_header_footer = True

def set_font(style, name, size, bold=False, color='000000'):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.append(fonts)
    for key in ['ascii','hAnsi','eastAsia','cs']:
        fonts.set(qn('w:'+key), name)

set_font(doc.styles['Normal'], 'Arial', 11)
normal = doc.styles['Normal'].paragraph_format
normal.space_after = Pt(7)
normal.line_spacing = 1.07
normal.widow_control = True
for style, size in [('Title',36),('Subtitle',19),('Heading 1',21),('Heading 2',12.5),('Heading 3',11)]:
    set_font(doc.styles[style], 'Arial',size, style != 'Subtitle')
    fmt = doc.styles[style].paragraph_format
    fmt.space_before = Pt(12 if style=='Heading 2' else 0)
    fmt.space_after = Pt(8 if style=='Heading 1' else 6)
    fmt.keep_with_next = True
    fmt.line_spacing = 1.02

for name,size,color in [('Caption',9,'465363'),('Header',8,'000000'),('Footer',8,'000000')]:
    set_font(doc.styles[name], 'Arial', size, color=color)
    doc.styles[name].paragraph_format.space_after = Pt(0)
set_font(doc.styles['List Bullet'],'Arial',11)
doc.styles['List Bullet'].paragraph_format.space_after=Pt(5)
doc.styles['List Bullet'].paragraph_format.line_spacing=1.04
doc.styles['List Bullet'].paragraph_format.left_indent=Inches(.17)
doc.styles['List Bullet'].paragraph_format.first_line_indent=Inches(-.12)

for name in ['Code','Meta','Numbered','TOC 1']:
    if name not in doc.styles: doc.styles.add_style(name,1)
set_font(doc.styles['Code'],'Consolas',9.3,color='18344A')
doc.styles['Code'].paragraph_format.space_after=Pt(0)
doc.styles['Code'].paragraph_format.line_spacing=Pt(11.4)
doc.styles['Code'].paragraph_format.keep_with_next=True
doc.styles['Code'].paragraph_format.left_indent=Inches(.08)
set_font(doc.styles['Meta'],'Arial',10,color='465363')
doc.styles['Meta'].paragraph_format.space_after=Pt(5)
set_font(doc.styles['Numbered'],'Arial',11)
doc.styles['Numbered'].paragraph_format.left_indent=Inches(.2)
doc.styles['Numbered'].paragraph_format.first_line_indent=Inches(-.2)
doc.styles['Numbered'].paragraph_format.space_after=Pt(5)
doc.styles['Numbered'].paragraph_format.line_spacing=1.04
set_font(doc.styles['TOC 1'],'Arial',10.5)
doc.styles['TOC 1'].paragraph_format.space_after=Pt(6)
doc.styles['TOC 1'].paragraph_format.line_spacing=1.05
doc.styles['TOC 1'].paragraph_format.tab_stops.add_tab_stop(Inches(6.9),WD_ALIGN_PARAGRAPH.RIGHT)

# Remove theme styling and any inherited title rules.
for st in doc.styles:
    if st.type == 1 and (st.name.startswith('Heading') or st.name in ['Title','Subtitle','Header']):
        for c in st.element.xpath('.//w:color'):
            for a in ['themeColor','themeTint','themeShade']:
                c.attrib.pop(qn('w:'+a),None)
            c.set(qn('w:val'),'000000')
        for border in st.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)

def inline(p,text):
    for token in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        if not token: continue
        r=p.add_run(token[2:-2] if token.startswith('**') else token[1:-1] if token.startswith('`') else token)
        if token.startswith('**'): r.bold=True
        if token.startswith('`'):
            r.font.name='Consolas'; r.font.size=Pt(10)
    return p

def field(p,instr,display=''):
    r=p.add_run(); begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin'); r._r.append(begin)
    r=p.add_run(); code=OxmlElement('w:instrText'); code.set(qn('xml:space'),'preserve'); code.text=instr; r._r.append(code)
    r=p.add_run(); sep=OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate'); r._r.append(sep)
    if display: p.add_run(display)
    r=p.add_run(); end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end'); r._r.append(end)

h=sec.header.paragraphs[0]
h.text='APEXFORMZ V2  /  PRODUCT REQUIREMENTS'
h.style=doc.styles['Header']
f=sec.footer.paragraphs[0]
f.style=doc.styles['Footer']
f.paragraph_format.tab_stops.clear_all()
f.paragraph_format.tab_stops.add_tab_stop(Inches(7),WD_ALIGN_PARAGRAPH.RIGHT)
f.add_run('VERSION 2.0  •  SEPTEMBER 2026\t')
field(f,' PAGE ')
f.add_run(' / ')
field(f,' NUMPAGES ')

def borders(table):
    pr=table._tbl.tblPr
    b=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        v=OxmlElement('w:'+side); v.set(qn('w:val'),'single'); v.set(qn('w:sz'),'4'); v.set(qn('w:color'),'D9D9D9'); b.append(v)
    pr.append(b)
    margins=OxmlElement('w:tblCellMar')
    for side,val in [('top',85),('bottom',85),('left',105),('right',105)]:
        m=OxmlElement('w:'+side);m.set(qn('w:w'),str(val));m.set(qn('w:type'),'dxa');margins.append(m)
    pr.append(margins)

def table(rows):
    n=len(rows[0])
    first=rows[0][0].lower()
    if n==2:
        widths=[1.7,5.3]
        if 'method' in first: widths=[3.1,3.9]
        if 'evidence' in first: widths=[2.7,4.3]
        if first=='term': widths=[1.45,5.55]
    elif n==3:
        if first=='id':
            widths=[.58,4.34,2.08]
            if rows[0][1]=='User need': widths=[.58,2.65,3.77]
        elif first=='priority': widths=[.62,2.04,4.34]
        elif first=='risk': widths=[1.4,2.3,3.3]
        else: widths=[1.3,3.15,2.55]
    else: widths=[7/n]*n
    t=doc.add_table(rows=0,cols=n)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    borders(t)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        trpr=t.rows[-1]._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit'); trpr.append(cant)
        if i==0:
            rep=OxmlElement('w:tblHeader'); trpr.append(rep)
        for j,(c,txt,w) in enumerate(zip(cells,row,widths)):
            c.width=Inches(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'18344A' if i==0 else ('F0F5F7' if i%2==0 else 'FFFFFF'));c._tc.get_or_add_tcPr().append(shade)
            p=c.paragraphs[0]
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.03
            p.paragraph_format.keep_with_next=(i==0)
            p.paragraph_format.widow_control=True
            if (j==0 and first in ['id','priority']):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            inline(p,txt)
            for r in p.runs:
                r.font.size=Pt(9.6)
                if i==0:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=Pt(2);p.add_run().font.size=Pt(2)
    return t

def image_block(name,caption):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.space_before=Pt(3)
    p.paragraph_format.space_after=Pt(5)
    sh=p.add_run().add_picture(str(ROOT/'figures'/name),width=Inches(6.9))
    sh._inline.docPr.set('descr',caption)
    c=doc.add_paragraph(caption,'Caption')
    c.paragraph_format.space_after=Pt(10)

# Cover uses the native Title style and no decorative rules.
p=doc.add_paragraph('APEXFORMZ V2','Title')
p.paragraph_format.space_before=Pt(44)
p.paragraph_format.space_after=Pt(10)
p=doc.add_paragraph('Product Requirements Document','Subtitle')
p.paragraph_format.space_after=Pt(22)
p=doc.add_paragraph('Modular human motion intelligence\nand sport expert engines')
p.paragraph_format.space_after=Pt(25)
for r in p.runs:r.font.size=Pt(16)
inline(doc.add_paragraph(), 'A shared motion foundation for independently evolving sport and movement analysis, real-time coaching, and athlete progression.')
p=doc.add_paragraph('Foundation release focus','Heading 2');p.paragraph_format.space_before=Pt(25)
inline(doc.add_paragraph(), 'Migrate the existing Taekwondo pipeline into isolated engines, preserve responsive feedback, and establish trustworthy scoring, reference data, and versioned results.')
table([
    ['Document control','Value'],
    ['Product','ApexFormz / SmartCoach'],
    ['Version and date','2.0  |  September 2026'],
    ['Document status','Master Product and Engineering PRD'],
    ['Review audience','Product, engineering, data, quality, and coaching'],
    ['Product and technical owners','To be assigned'],
    ['Release date and approvers','To be assigned'],
])
p=doc.add_paragraph('Decision required','Heading 2');p.paragraph_format.space_before=Pt(14)
inline(doc.add_paragraph(),'Confirm foundation scope, assign ownership, and resolve the acceptance and contract decisions listed in Section 20. Proposed additions are identified throughout the document.')

doc.add_page_break()
p=doc.add_paragraph('Contents','Title')
for r in p.runs:r.font.size=Pt(24)
p.paragraph_format.space_after=Pt(14)
p=doc.add_paragraph()
field(p,' TOC \\o "1-1" \\h \\z \\u ','Update this table of contents in Word.')
doc.add_paragraph('Reading guide','Heading 2')
inline(doc.add_paragraph(),'Sections 01–05 establish the product and architecture. Sections 06–16 define requirements and operations. Sections 17–20 cover migration, validation, acceptance, and open decisions. Appendices retain engineering contracts and illustrative configurations.')

text=(ROOT/'prd_content.md').read_text(encoding='utf-8')
pages=text.split('===PAGE===')
headings=[]
for page in pages:
    lines=page.strip().splitlines()
    idx=0
    while idx<len(lines):
        line=lines[idx].rstrip()
        if not line:
            idx+=1;continue
        if line.startswith('```'):
            idx+=1;block=[]
            while idx<len(lines) and not lines[idx].startswith('```'):
                block.append(lines[idx]);idx+=1
            for k,code in enumerate(block):
                p=doc.add_paragraph(code,'Code')
                if k==len(block)-1:
                    p.paragraph_format.keep_with_next=False
                    p.paragraph_format.space_after=Pt(9)
            idx+=1;continue
        if line.startswith('|'):
            rows=[]
            while idx<len(lines) and lines[idx].startswith('|'):
                row=[x.strip() for x in lines[idx].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[-: ]+',x) for x in row):rows.append(row)
                idx+=1
            table(rows);continue
        if line.startswith('!['):
            name,cap=line[2:-1].split('|',1)
            image_block(name,cap);idx+=1;continue
        if line.startswith('# '):
            title=line[2:];headings.append(title)
            p=doc.add_paragraph(title,'Heading 1')
            p.paragraph_format.page_break_before=True
        elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 2')
        elif line.startswith('- '):inline(doc.add_paragraph(style='List Bullet'),line[2:])
        elif re.match(r'^\d+\. ',line):inline(doc.add_paragraph(style='Numbered'),line)
        else:inline(doc.add_paragraph(),line)
        idx+=1

# Document metadata and automatic field refresh.
doc.core_properties.title='APEXFORMZ V2 Product Requirements Document'
doc.core_properties.subject='Modular human motion intelligence and sport expert engines'
doc.core_properties.author=''
doc.core_properties.keywords='ApexFormz, PRD, human motion, sport engines, V2'
doc.core_properties.comments=''
setting=OxmlElement('w:updateFields');setting.set(qn('w:val'),'true');doc.settings.element.append(setting)
doc.save(OUT)
print(str(OUT))
print(f'{len(headings)} main sections; {len(doc.tables)} tables; {len(doc.inline_shapes)} figures; {sum(len(p.text.split()) for p in doc.paragraphs)} prose words')
